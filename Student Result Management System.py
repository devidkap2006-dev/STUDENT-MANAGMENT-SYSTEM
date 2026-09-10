
"""
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# =====================================================
# EXCEL FILE SETUP
# =====================================================

FILE_NAME = "student_results.xlsx"

columns = [
    "Name", "Roll No.", "Class",
    "Subject 1", "Subject 2", "Subject 3",
    "Subject 4", "Subject 5",
    "Total Marks", "Percentage", "Result"
]

if not os.path.exists(FILE_NAME):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(columns)
    wb.save(FILE_NAME)


# =====================================================
# MAIN WINDOW
# =====================================================

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("900x600")
root.resizable(False, False)


# =====================================================
# MAIN FRAME
# =====================================================

main_frame = tk.Frame(root)
main_frame.pack(fill="both", expand=True, padx=20, pady=20)


# =====================================================
# CLEAR SCREEN
# =====================================================

def clear_screen():
    for widget in main_frame.winfo_children():
        widget.destroy()


# =====================================================
# TITLE
# =====================================================

def create_title(title):
    tk.Label(
        main_frame,
        text=title,
        font=("Arial", 24, "bold")
    ).pack(pady=20)


# =====================================================
# #1 ADD STUDENT
# =====================================================

def add_student():

    clear_screen()

    create_title("1. Add Student")

    form = tk.Frame(main_frame)
    form.pack(pady=10)

    labels = [
        "Student Name",
        "Roll No.",
        "Class",
        "Subject 1 Marks",
        "Subject 2 Marks",
        "Subject 3 Marks",
        "Subject 4 Marks",
        "Subject 5 Marks"
    ]

    entries = []

    for i, label in enumerate(labels):

        tk.Label(
            form,
            text=label + ":",
            font=("Arial", 11)
        ).grid(
            row=i,
            column=0,
            padx=15,
            pady=6,
            sticky="w"
        )

        entry = tk.Entry(
            form,
            width=30,
            font=("Arial", 11)
        )

        entry.grid(
            row=i,
            column=1,
            padx=15,
            pady=6
        )

        entries.append(entry)

    # -------------------------------------------------
    # SAVE FUNCTION
    # -------------------------------------------------

    def save_student():

        name = entries[0].get().strip()
        roll = entries[1].get().strip()
        student_class = entries[2].get().strip()

        # Check empty fields
        if name == "" or roll == "" or student_class == "":
            messagebox.showerror(
                "Error",
                "Please enter Name, Roll No. and Class."
            )
            return

        # Roll number validation
        try:
            roll_number = int(roll)
        except ValueError:
            messagebox.showerror(
                "Error",
                "Roll No. must be a number."
            )
            return

        # Marks validation
        marks = []

        try:
            for i in range(3, 8):

                mark = float(entries[i].get())

                if mark < 0 or mark > 100:
                    messagebox.showerror(
                        "Error",
                        "Marks must be between 0 and 100."
                    )
                    return

                marks.append(mark)

        except ValueError:
            messagebox.showerror(
                "Error",
                "Please enter valid marks."
            )
            return

        # -------------------------------------------------
        # OPEN EXCEL
        # -------------------------------------------------

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        # Check duplicate roll number
        for row in ws.iter_rows(min_row=2, values_only=True):

            if row[1] == roll_number:

                wb.close()

                messagebox.showerror(
                    "Error",
                    "This Roll No. already exists."
                )

                return

        # -------------------------------------------------
        # CALCULATIONS
        # -------------------------------------------------

        total = sum(marks)
        percentage = total / 5

        # Pass if every subject >= 40
        if all(mark >= 40 for mark in marks):
            result = "Pass"
        else:
            result = "Fail"

        # -------------------------------------------------
        # SAVE TO EXCEL
        # -------------------------------------------------

        ws.append([
            name,
            roll_number,
            student_class,
            *marks,
            total,
            percentage,
            result
        ])

        wb.save(FILE_NAME)
        wb.close()

        messagebox.showinfo(
            "Success",
            "Student record saved successfully!"
        )

        # IMPORTANT:
        # After Save → automatically open #2
        get_result()

    # -------------------------------------------------
    # SAVE BUTTON
    # -------------------------------------------------

    tk.Button(
        main_frame,
        text="💾 SAVE",
        font=("Arial", 13, "bold"),
        width=18,
        command=save_student
    ).pack(pady=20)


# =====================================================
# #2 GET RESULT
# =====================================================

def get_result():

    clear_screen()

    create_title("2. Get Result")

    tk.Label(
        main_frame,
        text="Enter Roll No.:",
        font=("Arial", 13)
    ).pack(pady=10)

    roll_entry = tk.Entry(
        main_frame,
        width=25,
        font=("Arial", 13)
    )

    roll_entry.pack(pady=5)

    result_frame = tk.Frame(main_frame)
    result_frame.pack(pady=20)

    # -------------------------------------------------
    # GET RESULT FUNCTION
    # -------------------------------------------------

    def search_result():

        roll = roll_entry.get().strip()

        if roll == "":
            messagebox.showerror(
                "Error",
                "Please enter Roll No."
            )
            return

        try:
            roll_number = int(roll)
        except ValueError:
            messagebox.showerror(
                "Error",
                "Roll No. must be a number."
            )
            return

        # Clear previous result
        for widget in result_frame.winfo_children():
            widget.destroy()

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        found = False

        for row in ws.iter_rows(min_row=2, values_only=True):

            if row[1] == roll_number:

                found = True

                # -------------------------------------------------
                # DISPLAY RESULT
                # -------------------------------------------------

                result_data = [
                    ("Name", row[0]),
                    ("Roll No.", row[1]),
                    ("Class", row[2]),
                    ("Total Marks", row[8]),
                    ("Percentage", f"{row[9]:.1f}%"),
                    ("Result", row[10])
                ]

                for i, (label, value) in enumerate(result_data):

                    tk.Label(
                        result_frame,
                        text=label + ":",
                        font=("Arial", 11, "bold"),
                        width=15,
                        anchor="w"
                    ).grid(
                        row=i,
                        column=0,
                        padx=10,
                        pady=5
                    )

                    tk.Label(
                        result_frame,
                        text=value,
                        font=("Arial", 11),
                        width=20,
                        anchor="w"
                    ).grid(
                        row=i,
                        column=1,
                        padx=10,
                        pady=5
                    )

                break

        wb.close()

        if not found:

            tk.Label(
                result_frame,
                text="❌ Student record not found.",
                font=("Arial", 13, "bold")
            ).pack(pady=20)

            return

        # -------------------------------------------------
        # AFTER SUCCESSFUL SUBMIT → #3
        # -------------------------------------------------

        tk.Button(
            main_frame,
            text="SUBMIT & CONTINUE",
            font=("Arial", 12, "bold"),
            width=22,
            command=show_all_results
        ).pack(pady=20)

    # -------------------------------------------------
    # GET RESULT BUTTON
    # -------------------------------------------------

    tk.Button(
        main_frame,
        text="🔍 GET RESULT",
        font=("Arial", 12, "bold"),
        width=18,
        command=search_result
    ).pack(pady=10)


# =====================================================
# #3 SHOW ALL RESULTS
# =====================================================

def show_all_results():

    clear_screen()

    create_title("3. Show All Results")

    # -------------------------------------------------
    # TREEVIEW
    # -------------------------------------------------

    table_frame = tk.Frame(main_frame)
    table_frame.pack(
        fill="both",
        expand=True,
        pady=10
    )

    display_columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )

    tree = ttk.Treeview(
        table_frame,
        columns=display_columns,
        show="headings",
        height=15
    )

    # Headings
    for column in display_columns:

        tree.heading(
            column,
            text=column
        )

        tree.column(
            column,
            width=130,
            anchor="center"
        )

    tree.pack(
        side="left",
        fill="both",
        expand=True
    )

    # Scrollbar
    scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

    tree.configure(
        yscrollcommand=scrollbar.set
    )

    scrollbar.pack(
        side="right",
        fill="y"
    )

    # -------------------------------------------------
    # READ EXCEL
    # -------------------------------------------------

    wb = load_workbook(FILE_NAME)
    ws = wb.active

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is not None:

            tree.insert(
                "",
                "end",
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[8],
                    f"{row[9]:.1f}%",
                    row[10]
                )
            )

    wb.close()


# =====================================================
# START PROGRAM
# =====================================================

add_student()

root.mainloop()

"""



import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# =====================================================
# EXCEL FILE SETUP
# =====================================================

FILE_NAME = "student_results.xlsx"

columns = [
    "Name", "Roll No.", "Class",
    "Subject 1", "Subject 2", "Subject 3",
    "Subject 4", "Subject 5",
    "Total Marks", "Percentage", "Result"
]

# Create Excel file if it does not exist
if not os.path.exists(FILE_NAME):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(columns)
    wb.save(FILE_NAME)


# =====================================================
# MAIN WINDOW
# =====================================================

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("900x600")
root.resizable(False, False)


# =====================================================
# MAIN FRAME
# =====================================================

main_frame = tk.Frame(root)
main_frame.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=20
)


# =====================================================
# CLEAR SCREEN
# =====================================================

def clear_screen():

    for widget in main_frame.winfo_children():
        widget.destroy()


# =====================================================
# TITLE
# =====================================================

def create_title(title):

    tk.Label(
        main_frame,
        text=title,
        font=("Arial", 24, "bold")
    ).pack(pady=20)


# =====================================================
# #1 ADD STUDENT
# =====================================================

def add_student():

    clear_screen()

    create_title("1. Add Student")

    form = tk.Frame(main_frame)
    form.pack(pady=10)

    labels = [
        "Student Name",
        "Roll No.",
        "Class",
        "Subject 1 Marks",
        "Subject 2 Marks",
        "Subject 3 Marks",
        "Subject 4 Marks",
        "Subject 5 Marks"
    ]

    entries = []

    # Create input fields
    for i, label in enumerate(labels):

        tk.Label(
            form,
            text=label + ":",
            font=("Arial", 11)
        ).grid(
            row=i,
            column=0,
            padx=15,
            pady=6,
            sticky="w"
        )

        entry = tk.Entry(
            form,
            width=30,
            font=("Arial", 11)
        )

        entry.grid(
            row=i,
            column=1,
            padx=15,
            pady=6
        )

        entries.append(entry)

    # -------------------------------------------------
    # SAVE STUDENT
    # -------------------------------------------------

    def save_student():

        name = entries[0].get().strip()
        roll = entries[1].get().strip()
        student_class = entries[2].get().strip()

        # Check basic details
        if name == "" or roll == "" or student_class == "":
            messagebox.showerror(
                "Error",
                "Please enter Name, Roll No. and Class."
            )
            return

        # Validate Roll No.
        try:
            roll_number = int(roll)

        except ValueError:
            messagebox.showerror(
                "Error",
                "Roll No. must be a number."
            )
            return

        # Validate marks
        marks = []

        try:

            for i in range(3, 8):

                mark = float(entries[i].get())

                if mark < 0 or mark > 100:

                    messagebox.showerror(
                        "Error",
                        "Marks must be between 0 and 100."
                    )

                    return

                marks.append(mark)

        except ValueError:

            messagebox.showerror(
                "Error",
                "Please enter valid marks."
            )

            return

        # -------------------------------------------------
        # OPEN EXCEL
        # -------------------------------------------------

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        # Check duplicate Roll No.
        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if row[1] == roll_number:

                wb.close()

                messagebox.showerror(
                    "Error",
                    "This Roll No. already exists."
                )

                return

        # -------------------------------------------------
        # CALCULATE TOTAL
        # -------------------------------------------------

        total = sum(marks)

        # Calculate percentage
        percentage = total / 5

        # -------------------------------------------------
        # CALCULATE RESULT
        # -------------------------------------------------

        # Pass if student gets 40 or more in every subject
        if all(mark >= 40 for mark in marks):
            result = "Pass"

        else:
            result = "Fail"

        # -------------------------------------------------
        # SAVE DATA TO EXCEL
        # -------------------------------------------------

        ws.append([
            name,
            roll_number,
            student_class,
            *marks,
            total,
            percentage,
            result
        ])

        wb.save(FILE_NAME)
        wb.close()

        messagebox.showinfo(
            "Success",
            "Student record saved successfully!"
        )

        # -------------------------------------------------
        # GO TO #2 AUTOMATICALLY
        # -------------------------------------------------

        get_result()

    # -------------------------------------------------
    # SAVE BUTTON
    # -------------------------------------------------

    tk.Button(
        main_frame,
        text="💾 SAVE",
        font=("Arial", 13, "bold"),
        width=18,
        command=save_student
    ).pack(pady=20)


# =====================================================
# #2 GET RESULT
# =====================================================

def get_result():

    clear_screen()

    create_title("2. Get Result")

    tk.Label(
        main_frame,
        text="Enter Roll No.:",
        font=("Arial", 13)
    ).pack(pady=10)

    roll_entry = tk.Entry(
        main_frame,
        width=25,
        font=("Arial", 13)
    )

    roll_entry.pack(pady=5)

    result_frame = tk.Frame(main_frame)
    result_frame.pack(pady=20)

    # -------------------------------------------------
    # SEARCH RESULT
    # -------------------------------------------------

    def search_result():

        roll = roll_entry.get().strip()

        if roll == "":

            messagebox.showerror(
                "Error",
                "Please enter Roll No."
            )

            return

        try:

            roll_number = int(roll)

        except ValueError:

            messagebox.showerror(
                "Error",
                "Roll No. must be a number."
            )

            return

        # Clear previous result
        for widget in result_frame.winfo_children():
            widget.destroy()

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        found = False

        # Search Roll No.
        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if row[1] == roll_number:

                found = True

                # Result information
                result_data = [
                    ("Name", row[0]),
                    ("Roll No.", row[1]),
                    ("Class", row[2]),
                    ("Total Marks", row[8]),
                    ("Percentage", f"{row[9]:.1f}%"),
                    ("Result", row[10])
                ]

                # Display result
                for i, (label, value) in enumerate(result_data):

                    tk.Label(
                        result_frame,
                        text=label + ":",
                        font=("Arial", 11, "bold"),
                        width=15,
                        anchor="w"
                    ).grid(
                        row=i,
                        column=0,
                        padx=10,
                        pady=5
                    )

                    tk.Label(
                        result_frame,
                        text=value,
                        font=("Arial", 11),
                        width=20,
                        anchor="w"
                    ).grid(
                        row=i,
                        column=1,
                        padx=10,
                        pady=5
                    )

                break

        wb.close()

        # -------------------------------------------------
        # IF RECORD NOT FOUND
        # -------------------------------------------------

        if not found:

            tk.Label(
                result_frame,
                text="❌ Student record not found.",
                font=("Arial", 13, "bold")
            ).pack(pady=20)

            return

        # -------------------------------------------------
        # SUBMIT & CONTINUE BUTTON
        # -------------------------------------------------

        tk.Button(
            main_frame,
            text="SUBMIT & CONTINUE",
            font=("Arial", 12, "bold"),
            width=22,
            command=show_all_results
        ).pack(pady=20)

    # -------------------------------------------------
    # GET RESULT BUTTON
    # -------------------------------------------------

    tk.Button(
        main_frame,
        text="🔍 GET RESULT",
        font=("Arial", 12, "bold"),
        width=18,
        command=search_result
    ).pack(pady=10)


# =====================================================
# #3 SHOW ALL RESULTS
# =====================================================

def show_all_results():

    clear_screen()

    create_title("3. Show All Results")

    # -------------------------------------------------
    # TABLE FRAME
    # -------------------------------------------------

    table_frame = tk.Frame(main_frame)

    table_frame.pack(
        fill="both",
        expand=True,
        pady=10
    )

    display_columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )

    # -------------------------------------------------
    # TREEVIEW
    # -------------------------------------------------

    tree = ttk.Treeview(
        table_frame,
        columns=display_columns,
        show="headings",
        height=14
    )

    # Set headings and columns
    for column in display_columns:

        tree.heading(
            column,
            text=column
        )

        tree.column(
            column,
            width=130,
            anchor="center"
        )

    tree.pack(
        side="left",
        fill="both",
        expand=True
    )

    # -------------------------------------------------
    # SCROLLBAR
    # -------------------------------------------------

    scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

    tree.configure(
        yscrollcommand=scrollbar.set
    )

    scrollbar.pack(
        side="right",
        fill="y"
    )

    # -------------------------------------------------
    # READ DATA FROM EXCEL
    # -------------------------------------------------

    wb = load_workbook(FILE_NAME)
    ws = wb.active

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is not None:

            tree.insert(
                "",
                "end",
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[8],
                    f"{row[9]:.1f}%",
                    row[10]
                )
            )

    wb.close()

    # =================================================
    # BUTTONS
    # =================================================

    button_frame = tk.Frame(main_frame)
    button_frame.pack(pady=15)

    # -------------------------------------------------
    # ADD ANOTHER STUDENT
    # -------------------------------------------------

    tk.Button(
        button_frame,
        text="➕ ADD ANOTHER STUDENT",
        font=("Arial", 12, "bold"),
        width=25,
        command=add_student
    ).grid(
        row=0,
        column=0,
        padx=10
    )

    # -------------------------------------------------
    # EXIT
    # -------------------------------------------------

    tk.Button(
        button_frame,
        text="EXIT",
        font=("Arial", 12, "bold"),
        width=15,
        command=root.destroy
    ).grid(
        row=0,
        column=1,
        padx=10
    )


# =====================================================
# START PROGRAM
# =====================================================

# Program starts directly from #1
add_student()

root.mainloop()
